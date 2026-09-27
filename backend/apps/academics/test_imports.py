import io

from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from rest_framework import status

from .models import Classroom, Subject, TeachingAssignment
from .tests import AcademicsTestBase


def excel(sheet_title, header, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile("data.xlsx", buffer.getvalue())


class AdminImportTests(AcademicsTestBase):
    def setUp(self):
        super().setUp()
        self.client.force_authenticate(self.admin)

    def post(self, name, file, dry_run=False):
        return self.client.post(
            reverse(name), {"file": file, "dry_run": str(dry_run).lower()}, format="multipart"
        )

    def test_templates_download_with_dropdowns(self):
        for name, sheet in [
            ("subject-import-template", "Mapel"),
            ("classroom-import-template", "Kelas"),
            ("teaching-assignment-import-template", "Pengampu"),
        ]:
            res = self.client.get(reverse(name))
            self.assertEqual(res.status_code, status.HTTP_200_OK, name)
            workbook = load_workbook(io.BytesIO(res.content))
            self.assertEqual(workbook.active.title, sheet)
        # Dropdown pengampu berisi kelas, kode mapel, dan username guru yang ada.
        workbook = load_workbook(
            io.BytesIO(self.client.get(reverse("teaching-assignment-import-template")).content)
        )
        values = {v for row in workbook["Daftar"].iter_rows(values_only=True) for v in row if v}
        self.assertTrue({"XI RPL 1", "PWEB", "guru1"} <= values)

    def test_import_subjects(self):
        rows = [["mtk", "Matematika"], ["BARAB", "Bahasa Arab (اللغة العربية)"]]
        res = self.post("subject-import-rows", excel("Mapel", ["Kode*", "Nama Mapel*"], rows))
        self.assertEqual(res.status_code, status.HTTP_201_CREATED, res.data)
        self.assertEqual(Subject.objects.get(code="MTK").name, "Matematika")
        self.assertEqual(Subject.objects.get(code="BARAB").name, "Bahasa Arab (اللغة العربية)")

    def test_subject_errors_all_or_nothing(self):
        rows = [
            ["BIG", "Bahasa Inggris"],
            ["pweb", "Sudah ada"],
            ["big", "Ganda di file"],
            ["", "Tanpa kode"],
        ]
        res = self.post("subject-import-rows", excel("Mapel", ["Kode*", "Nama Mapel*"], rows), dry_run=True)
        self.assertEqual(res.data["summary"], {"total": 4, "valid": 1, "invalid": 3})
        errors = [" ".join(r["errors"]) for r in res.data["rows"]]
        self.assertIn("Kode", errors[1])
        self.assertIn("sama dengan baris 2", errors[2])
        res = self.post("subject-import-rows", excel("Mapel", ["Kode*", "Nama Mapel*"], rows))
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(Subject.objects.filter(code="BIG").exists())

    def test_import_classrooms_with_homeroom(self):
        header = ["Nama Kelas*", "Tingkat*", "Jurusan*", "Tahun Ajaran*", "Username Wali Kelas"]
        rows = [["X RPL 2", "X", "rpl", "2026/2027", "guru1"], ["XII TKJ 1", "XII", "TKJ", "2026/2027", ""]]
        res = self.post("classroom-import-rows", excel("Kelas", header, rows))
        self.assertEqual(res.status_code, status.HTTP_201_CREATED, res.data)
        classroom = Classroom.objects.get(name="X RPL 2")
        self.assertEqual((classroom.major, classroom.homeroom_teacher), ("RPL", self.teacher))

    def test_classroom_errors(self):
        header = ["Nama Kelas*", "Tingkat*", "Jurusan*", "Tahun Ajaran*", "Username Wali Kelas"]
        rows = [
            ["XI RPL 1", "XI", "RPL", "2026/2027", ""],  # sudah ada
            ["X A", "XIII", "RPL", "2026/2027", ""],  # tingkat salah
            ["X B", "X", "RPL", "2026-2027", ""],  # format tahun salah
            ["X C", "X", "RPL", "2026/2027", "tidak.ada"],  # guru tidak ada
            ["X D", "X", "RPL", "2026/2027", "siswa1"],  # bukan guru
        ]
        res = self.post("classroom-import-rows", excel("Kelas", header, rows), dry_run=True)
        self.assertEqual(res.data["summary"]["invalid"], 5, res.data)

    def test_import_assignments(self):
        header = ["Kelas*", "Tahun Ajaran", "Kode Mapel*", "Username Guru*"]
        rows = [["xi rpl 1", "", "jarkom", "guru2"], ["XI TKJ 1", "2026/2027", "PWEB", "guru1"]]
        res = self.post("teaching-assignment-import-rows", excel("Pengampu", header, rows))
        self.assertEqual(res.status_code, status.HTTP_201_CREATED, res.data)
        self.assertTrue(
            TeachingAssignment.objects.filter(
                classroom=self.rpl, subject=self.net, teacher=self.other_teacher
            ).exists()
        )

    def test_assignment_errors(self):
        header = ["Kelas*", "Tahun Ajaran", "Kode Mapel*", "Username Guru*"]
        rows = [
            ["XI RPL 1", "", "PWEB", "guru1"],  # sudah ada
            ["XII MM 9", "", "PWEB", "guru1"],  # kelas tidak ada
            ["XI RPL 1", "", "XYZ", "guru1"],  # mapel tidak ada
            ["XI RPL 1", "", "JARKOM", "siswa1"],  # bukan guru
        ]
        res = self.post("teaching-assignment-import-rows", excel("Pengampu", header, rows), dry_run=True)
        self.assertEqual(res.data["summary"]["invalid"], 4, res.data)
        errors = [" ".join(r["errors"]) for r in res.data["rows"]]
        self.assertIn("tidak ditemukan", errors[1])
        self.assertIn("XYZ", errors[2])

    def test_only_admin(self):
        self.client.force_authenticate(self.teacher)
        self.assertEqual(
            self.client.get(reverse("subject-import-template")).status_code, status.HTTP_403_FORBIDDEN
        )
        res = self.post("classroom-import-rows", excel("Kelas", ["Nama Kelas*"], [["X"]]))
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)
