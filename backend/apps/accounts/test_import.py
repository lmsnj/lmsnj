import io

from django.contrib.auth.hashers import identify_hasher
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from apps.academics.models import Classroom

from .models import Role, User
from .user_import import IMPORT_HASH_ITERATIONS

PASSWORD = "RahasiaKuat#2026"
HEADER = ["Nama Lengkap*", "Username", "Peran*", "NIS", "NIP", "Kelas", "Email", "Password"]


def excel(rows, header=HEADER, name="pengguna.xlsx", sheet_title="Pengguna"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(name, buffer.getvalue())


class UserImportTests(APITestCase):
    url = reverse("user-import-users")

    def setUp(self):
        self.admin = User.objects.create_user(
            username="admin1", password=PASSWORD, full_name="Admin", role=Role.ADMIN
        )
        self.rpl = Classroom.objects.create(
            name="XI RPL 1", grade="XI", major="RPL", academic_year="2026/2027"
        )
        Classroom.objects.create(name="XI RPL 1", grade="XI", major="RPL", academic_year="2025/2026")
        User.objects.create_user(
            username="rizki", password=PASSWORD, full_name="Rizki", role=Role.STUDENT, nis="26001"
        )
        self.client.force_authenticate(self.admin)

    def upload(self, file, dry_run=True):
        return self.client.post(self.url, {"file": file, "dry_run": str(dry_run).lower()}, format="multipart")

    def test_template_download(self):
        res = self.client.get(reverse("user-import-template"))
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertIn("Template Impor Pengguna.xlsx", res["Content-Disposition"])
        workbook = load_workbook(io.BytesIO(res.content))
        self.assertEqual([c.value for c in workbook["Pengguna"][1]], HEADER)
        # Dropdown kelas hanya berisi kelas tahun ajaran terbaru.
        self.assertEqual([r[0].value for r in workbook["DaftarKelas"].iter_rows()], ["XI RPL 1"])

    def test_template_can_be_uploaded_back_after_filling(self):
        res = self.client.get(reverse("user-import-template"))
        workbook = load_workbook(io.BytesIO(res.content))
        workbook["Pengguna"].append(["Nabila Zahra", "", "siswa", "26005", "", "XI RPL 1", "", ""])
        buffer = io.BytesIO()
        workbook.save(buffer)
        res = self.upload(SimpleUploadedFile("isi.xlsx", buffer.getvalue()))
        self.assertEqual(res.data["summary"], {"total": 1, "valid": 1, "invalid": 0}, res.data)

    def test_dry_run_saves_nothing(self):
        res = self.upload(excel([["Nabila Zahra", "nabila", "siswa", "26005", "", "XI RPL 1", "", ""]]))
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertFalse(res.data["imported"])
        self.assertEqual(res.data["rows"][0]["classroom"], "XI RPL 1")
        self.assertTrue(res.data["rows"][0]["password_generated"])
        self.assertFalse(User.objects.filter(username="nabila").exists())

    def test_import_creates_users_with_working_passwords(self):
        rows = [
            ["Nabila Zahra", "", "siswa", 26005, "", "xi  rpl 1", "", ""],  # NIS angka, kelas beda spasi
            ["Siti Aminah", "bu.siti", "Guru", "", "198702022012", "", "siti@smknj.sch.id", "BukuPena#2026"],
        ]
        res = self.upload(excel(rows), dry_run=False)
        self.assertEqual(res.status_code, status.HTTP_201_CREATED, res.data)
        accounts = {a["username"]: a for a in res.data["accounts"]}

        student = User.objects.get(username="26005")  # username diambil dari NIS
        self.assertEqual((student.role, student.nis, student.classroom), (Role.STUDENT, "26005", self.rpl))
        teacher = User.objects.get(username="bu.siti")
        self.assertEqual((teacher.role, teacher.nip, teacher.classroom), (Role.TEACHER, "198702022012", None))
        self.assertTrue(accounts["26005"]["generated"])
        self.assertFalse(accounts["bu.siti"]["generated"])

        # Password (otomatis maupun dari file) bisa dipakai login.
        self.client.force_authenticate(None)
        for username, account in accounts.items():
            res = self.client.post(
                reverse("auth-login"), {"username": username, "password": account["password"]}
            )
            self.assertEqual(res.status_code, status.HTTP_200_OK, username)

    def test_initial_hash_is_upgraded_on_first_login(self):
        rows = [["Nabila Zahra", "nabila", "siswa", "26005", "", "", "", "BukuPena#2026"]]
        self.upload(excel(rows), dry_run=False)
        user = User.objects.get(username="nabila")
        self.assertEqual(
            identify_hasher(user.password).decode(user.password)["iterations"], IMPORT_HASH_ITERATIONS
        )
        self.client.force_authenticate(None)
        self.client.post(reverse("auth-login"), {"username": "nabila", "password": "BukuPena#2026"})
        user.refresh_from_db()
        self.assertGreater(
            identify_hasher(user.password).decode(user.password)["iterations"], IMPORT_HASH_ITERATIONS
        )

    def test_row_errors_and_all_or_nothing(self):
        rows = [
            ["Baris Benar", "benar", "siswa", "26010", "", "", "", ""],
            ["", "tanpanama", "siswa", "", "", "", "", ""],
            ["Rizki Kembar", "RIZKI", "siswa", "26001", "", "", "", ""],  # username & NIS sudah ada
            ["Salah Peran", "x1", "kepsek", "", "", "", "", ""],
            ["Kelas Salah", "x2", "siswa", "", "", "XII RPL 9", "", ""],
            ["Ada Spasi", "ada spasi", "guru", "", "", "", "", ""],
            ["Email Salah", "x3", "guru", "", "", "", "bukan-email", ""],
            ["Password Lemah", "x4", "guru", "", "", "", "", "123"],
            ["Ganda", "benar", "guru", "", "", "", "", ""],  # username sama dengan baris 2
            ["Tanpa Username", "", "guru", "", "", "", "", ""],
        ]
        res = self.upload(excel(rows))
        self.assertEqual(res.data["summary"], {"total": 10, "valid": 1, "invalid": 9})
        errors = {r["row"]: " ".join(r["errors"]) for r in res.data["rows"]}
        self.assertIn("Nama lengkap wajib", errors[3])
        self.assertIn("Username 'RIZKI' sudah dipakai", errors[4])
        self.assertIn("NIS 26001 sudah dipakai", errors[4])
        self.assertIn("Peran harus", errors[5])
        self.assertIn("Kelas 'XII RPL 9' tidak ditemukan di tahun ajaran 2026/2027", errors[6])
        self.assertIn("tanpa spasi", errors[7])
        self.assertIn("email", errors[8])
        self.assertIn("Password", errors[9])
        self.assertIn("sama dengan baris 2", errors[10])
        self.assertIn("Username wajib", errors[11])

        # Tidak ada yang disimpan selama masih ada baris yang salah.
        res = self.upload(excel(rows), dry_run=False)
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(User.objects.filter(username="benar").exists())

    def test_nis_ignored_for_teacher_and_blank_rows_skipped(self):
        rows = [["Guru Baru", "guru.baru", "guru", "26001", "", "XI RPL 1", "", ""], [None] * 8]
        res = self.upload(excel(rows), dry_run=False)
        self.assertEqual(res.status_code, status.HTTP_201_CREATED, res.data)
        guru = User.objects.get(username="guru.baru")
        self.assertEqual((guru.nis, guru.classroom), (None, None))
        self.assertEqual(res.data["summary"]["total"], 1)

    def test_bad_files(self):
        res = self.upload(SimpleUploadedFile("data.xls", b"lama"))
        self.assertIn("file", res.data)
        res = self.upload(SimpleUploadedFile("rusak.xlsx", b"bukan excel"))
        self.assertIn("tidak bisa dibaca", res.data["file"])
        res = self.upload(excel([["A", "a", "siswa"]], header=["Nama", "Username", "Peran"]))
        self.assertIn("Nama Lengkap", res.data["file"])
        res = self.upload(excel([]))
        self.assertIn("Tidak ada data", res.data["file"])
        self.assertEqual(self.client.post(self.url, {}).status_code, status.HTTP_400_BAD_REQUEST)

    def test_only_admin(self):
        teacher = User.objects.create_user(username="g", password=PASSWORD, full_name="G", role=Role.TEACHER)
        self.client.force_authenticate(teacher)
        self.assertEqual(
            self.client.get(reverse("user-import-template")).status_code, status.HTTP_403_FORBIDDEN
        )
        self.assertEqual(self.upload(excel([])).status_code, status.HTTP_403_FORBIDDEN)
