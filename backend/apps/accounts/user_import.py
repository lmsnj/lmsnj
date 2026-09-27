"""Impor pengguna dari file Excel (.xlsx).

Alur: admin mengunduh template -> mengisi -> mengunggah untuk DIPERIKSA (tidak ada yang
disimpan) -> jika semua baris benar, mengunggah lagi untuk DIIMPOR. Impor bersifat semua
atau tidak sama sekali, agar tidak ada data setengah jadi saat file diperbaiki lalu diunggah ulang.
"""

import io
import os
import secrets
from concurrent.futures import ThreadPoolExecutor

from django.contrib.auth.hashers import get_hasher
from django.contrib.auth.password_validation import validate_password
from django.contrib.auth.validators import UnicodeUsernameValidator
from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.validators import validate_email
from django.db import IntegrityError, transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from rest_framework.exceptions import ValidationError

from apps.academics.models import Classroom

from .models import Role, User

MAX_ROWS = 1000
DATA_SHEET = "Pengguna"
CLASS_LIST_SHEET = "DaftarKelas"

# (kunci, judul kolom, wajib)
COLUMNS = [
    ("full_name", "Nama Lengkap", True),
    ("username", "Username", False),
    ("role", "Peran", True),
    ("nis", "NIS", False),
    ("nip", "NIP", False),
    ("classroom", "Kelas", False),
    ("email", "Email", False),
    ("password", "Password", False),
]
TEXT_COLUMNS = {"username", "nis", "nip", "password"}  # diformat teks agar nol di depan tidak hilang

ROLE_ALIASES = {
    "siswa": Role.STUDENT,
    "student": Role.STUDENT,
    "guru": Role.TEACHER,
    "teacher": Role.TEACHER,
    "admin": Role.ADMIN,
}

# Password awal otomatis: kata sederhana + 4 angka, mudah diketik siswa (contoh: pohon4821).
PASSWORD_WORDS = [
    "awan", "bintang", "buku", "bulan", "danau", "embun", "gunung", "hujan", "kopi", "laut",
    "meja", "mentari", "pelangi", "pena", "pohon", "rimba", "samudra", "sungai", "taman", "angin",
]  # fmt: skip

# Password awal di-hash dengan iterasi lebih ringan agar impor ratusan akun tidak melewati batas
# waktu request. Aman karena bersifat sementara: saat pengguna login pertama kali, Django otomatis
# meng-hash ulang dengan iterasi standar (PBKDF2PasswordHasher.must_update).
IMPORT_HASH_ITERATIONS = 100_000

username_validator = UnicodeUsernameValidator()


def generate_password():
    return f"{secrets.choice(PASSWORD_WORDS)}{secrets.randbelow(9000) + 1000}"


def hash_initial_password(password):
    hasher = get_hasher("pbkdf2_sha256")
    return hasher.encode(password, hasher.salt(), iterations=IMPORT_HASH_ITERATIONS)


def normalize(text):
    return " ".join(str(text).split()).lower()


def latest_classrooms():
    """Kelas di tahun ajaran terbaru, dicari berdasarkan nama (tidak peka huruf besar/spasi)."""
    year = Classroom.objects.order_by("-academic_year").values_list("academic_year", flat=True).first()
    classrooms = Classroom.objects.filter(academic_year=year).order_by("grade", "name")
    return year, {normalize(c.name): c for c in classrooms}


# ---------------------------------------------------------------- template


def build_template():
    year, classrooms = latest_classrooms()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = DATA_SHEET

    header_fill = PatternFill("solid", fgColor="059669")
    for index, (key, label, required) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=f"{label}*" if required else label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        letter = cell.column_letter
        sheet.column_dimensions[letter].width = 28 if key in ("full_name", "email") else 16
        if key in TEXT_COLUMNS:
            for row in range(2, MAX_ROWS + 2):
                sheet[f"{letter}{row}"].number_format = "@"
    sheet.freeze_panes = "A2"

    last_row = MAX_ROWS + 1
    role_validation = DataValidation(type="list", formula1='"siswa,guru,admin"', allow_blank=False)
    role_validation.error = "Pilih siswa, guru, atau admin."
    sheet.add_data_validation(role_validation)
    role_validation.add(f"C2:C{last_row}")

    # Daftar kelas di lembar tersembunyi, dipakai sebagai pilihan dropdown kolom Kelas.
    if classrooms:
        class_sheet = workbook.create_sheet(CLASS_LIST_SHEET)
        for row, classroom in enumerate(classrooms.values(), start=1):
            class_sheet.cell(row=row, column=1, value=classroom.name)
        class_sheet.sheet_state = "hidden"
        class_validation = DataValidation(
            type="list", formula1=f"{CLASS_LIST_SHEET}!$A$1:$A${len(classrooms)}", allow_blank=True
        )
        class_validation.errorStyle = "warning"
        class_validation.error = "Kelas tidak ada di daftar. Periksa kembali nama kelasnya."
        sheet.add_data_validation(class_validation)
        class_validation.add(f"F2:F{last_row}")

    guide = workbook.create_sheet("Petunjuk", 0)
    lines = [
        ("Petunjuk Impor Pengguna — LMS SMKS Nurul Jadid", True),
        ("", False),
        (f"Isi data di lembar '{DATA_SHEET}', mulai baris ke-2. Maksimal {MAX_ROWS} baris per file.", False),
        ("Kolom bertanda * wajib diisi.", False),
        ("", False),
        ("Nama Lengkap* : nama siswa/guru.", False),
        ("Username      : untuk login, tanpa spasi. Siswa boleh dikosongkan jika NIS diisi (NIS jadi username).", False),
        ("Peran*        : siswa, guru, atau admin (pilih dari daftar).", False),
        ("NIS           : khusus siswa. NIP: khusus guru. Keduanya tidak boleh sama dengan akun lain.", False),
        (
            f"Kelas         : khusus siswa, pilih dari daftar (tahun ajaran {year})." if year
            else "Kelas         : belum ada kelas. Buat kelas di menu Kelas terlebih dahulu, atau kosongkan.",
            False,
        ),
        ("Email         : opsional.", False),
        ("Password      : opsional. Jika kosong, dibuat otomatis (contoh: pohon4821) dan bisa diunduh setelah impor.", False),
        ("", False),
        ("Contoh isian:", True),
    ]  # fmt: skip
    for row, (text, bold) in enumerate(lines, start=1):
        guide.cell(row=row, column=1, value=text).font = Font(bold=bold, size=13 if row == 1 else 11)
    example_class = next(iter(classrooms.values())).name if classrooms else ""
    examples = [
        [label for _, label, _ in COLUMNS],
        ["Muhammad Rizki", "", "siswa", "26001", "", example_class, "", ""],
        ["Siti Aminah", "bu.siti", "guru", "", "198702022012", "", "siti@smknj.sch.id", ""],
    ]
    start = len(lines) + 1
    for offset, values in enumerate(examples):
        for column, value in enumerate(values, start=1):
            cell = guide.cell(row=start + offset, column=column, value=value)
            cell.font = Font(bold=offset == 0)
    guide.column_dimensions["A"].width = 30
    # Excel langsung membuka lembar isian, bukan petunjuk.
    workbook.active = workbook.index(sheet)
    sheet.sheet_view.tabSelected = True
    guide.sheet_view.tabSelected = False

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------- membaca file


def cell_text(value):
    """Ubah isi sel menjadi teks. Angka seperti 26001 (terbaca 26001.0) menjadi '26001'."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def read_rows(file):
    if os.path.splitext(file.name)[1].lower() != ".xlsx":
        raise ValidationError(
            {"file": "Gunakan file Excel berformat .xlsx (unduh template yang disediakan)."}
        )
    try:
        workbook = load_workbook(file, read_only=True, data_only=True)
    except Exception:
        raise ValidationError({"file": "File tidak bisa dibaca. Pastikan file Excel .xlsx yang valid."})

    sheet = workbook[DATA_SHEET] if DATA_SHEET in workbook.sheetnames else workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None) or ()
    labels = {normalize(label): key for key, label, _ in COLUMNS}
    positions = {}
    for index, title in enumerate(header):
        key = labels.get(normalize(cell_text(title).replace("*", "")))
        if key:
            positions[key] = index
    missing = [label for key, label, required in COLUMNS if required and key not in positions]
    if missing:
        raise ValidationError(
            {"file": f"Kolom {', '.join(missing)} tidak ditemukan. Gunakan template yang disediakan."}
        )

    result = []
    for row_number, values in enumerate(rows, start=2):
        data = {
            key: cell_text(values[index]) if index < len(values) else "" for key, index in positions.items()
        }
        if not any(data.values()):
            continue  # baris kosong dilewati
        result.append({"row": row_number, "data": {key: data.get(key, "") for key, _, _ in COLUMNS}})
        if len(result) > MAX_ROWS:
            raise ValidationError(
                {"file": f"Maksimal {MAX_ROWS} pengguna per file. Bagi menjadi beberapa file."}
            )
    workbook.close()
    if not result:
        raise ValidationError({"file": f"Tidak ada data di lembar '{DATA_SHEET}'."})
    return result


# ---------------------------------------------------------------- validasi


def validate_rows(rows):
    """Periksa setiap baris. Mengembalikan baris yang sudah dibersihkan + daftar error per baris."""
    year, classrooms = latest_classrooms()
    # Dibandingkan tanpa peduli huruf besar agar "Rizki" dan "rizki" tidak menjadi dua akun.
    # Jumlah akun satu sekolah kecil, jadi cukup ambil semua username sekaligus.
    taken_usernames = {u.lower() for u in User.objects.values_list("username", flat=True)}
    taken_nis = set(User.objects.exclude(nis=None).values_list("nis", flat=True))
    taken_nip = set(User.objects.exclude(nip=None).values_list("nip", flat=True))
    seen = {"username": {}, "nis": {}, "nip": {}}

    results = []
    for row in rows:
        data = row["data"]
        errors = []
        role = ROLE_ALIASES.get(data["role"].lower())
        if not data["full_name"]:
            errors.append("Nama lengkap wajib diisi.")
        if role is None:
            errors.append("Peran harus siswa, guru, atau admin.")

        username = data["username"] or (data["nis"] if role == Role.STUDENT else "")
        if not username:
            errors.append("Username wajib diisi (siswa boleh dikosongkan jika NIS diisi).")
        else:
            try:
                username_validator(username)
            except DjangoValidationError:
                errors.append("Username hanya boleh huruf, angka, dan . @ + - _ (tanpa spasi).")
            if len(username) > 150:
                errors.append("Username terlalu panjang (maks. 150 karakter).")
            if username.lower() in taken_usernames:
                errors.append(f"Username '{username}' sudah dipakai akun lain.")

        nis = data["nis"] if role == Role.STUDENT else ""
        nip = data["nip"] if role == Role.TEACHER else ""
        if nis and nis in taken_nis:
            errors.append(f"NIS {nis} sudah dipakai akun lain.")
        if nip and nip in taken_nip:
            errors.append(f"NIP {nip} sudah dipakai akun lain.")

        classroom = None
        if role == Role.STUDENT and data["classroom"]:
            classroom = classrooms.get(normalize(data["classroom"]))
            if classroom is None:
                where = f"tahun ajaran {year}" if year else "sistem (belum ada kelas)"
                errors.append(f"Kelas '{data['classroom']}' tidak ditemukan di {where}.")

        if data["email"]:
            try:
                validate_email(data["email"])
            except DjangoValidationError:
                errors.append("Format email tidak valid.")

        if data["password"]:
            try:
                candidate = User(username=username, full_name=data["full_name"], email=data["email"])
                validate_password(data["password"], user=candidate)
            except DjangoValidationError as error:
                errors.append("Password: " + " ".join(error.messages))

        # Duplikat di dalam file yang sama.
        for field, label, value in (
            ("username", "Username", username.lower()),
            ("nis", "NIS", nis),
            ("nip", "NIP", nip),
        ):
            if not value:
                continue
            if value in seen[field]:
                errors.append(f"{label} sama dengan baris {seen[field][value]}.")
            else:
                seen[field][value] = row["row"]

        results.append(
            {
                "row": row["row"],
                "full_name": data["full_name"],
                "username": username,
                "role": role,
                "nis": nis or None,
                "nip": nip or None,
                "classroom": classroom,
                "email": data["email"],
                "password": data["password"],
                "errors": errors,
            }
        )
    return results


# ---------------------------------------------------------------- simpan


def create_users(results):
    """Simpan semua baris (dipanggil hanya jika tidak ada error). Mengembalikan daftar akun + password."""
    passwords = [r["password"] or generate_password() for r in results]
    # Hashing password lambat (disengaja demi keamanan); dikerjakan paralel.
    with ThreadPoolExecutor(max_workers=min(8, os.cpu_count() or 2)) as pool:
        hashes = list(pool.map(hash_initial_password, passwords))

    users = [
        User(
            username=r["username"],
            full_name=r["full_name"],
            role=r["role"],
            nis=r["nis"],
            nip=r["nip"],
            classroom=r["classroom"],
            email=r["email"],
            password=password_hash,
        )
        for r, password_hash in zip(results, hashes)
    ]
    try:
        with transaction.atomic():
            User.objects.bulk_create(users)
    except IntegrityError:
        # Sangat jarang: data yang sama dibuat orang lain di antara pemeriksaan dan impor.
        raise ValidationError(
            {"file": "Sebagian data baru saja dipakai akun lain. Periksa ulang file lalu coba lagi."}
        )

    return [
        {
            "row": r["row"],
            "full_name": r["full_name"],
            "username": r["username"],
            "role": r["role"],
            "classroom": r["classroom"].name if r["classroom"] else "",
            "password": password,
            "generated": not r["password"],
        }
        for r, password in zip(results, passwords)
    ]
