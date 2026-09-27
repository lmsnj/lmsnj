import io
import shutil
import tempfile
import unicodedata
import zipfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from rest_framework import status

from .models import Question, QuestionImage
from .question_import import COLUMNS, parse_text
from .tests import AssessmentTestBase

TEMP_MEDIA = tempfile.mkdtemp()
# PNG minimal: tanda tangan PNG yang valid + isi tiruan (cukup untuk pemeriksaan tanda tangan).
PNG = bytes.fromhex("89504e470d0a1a0a") + b"0" * 32
JPG = bytes.fromhex("ffd8ffe0") + b"0" * 32
RLM, LRM, NBSP = chr(0x200F), chr(0x200E), chr(0xA0)

# Teks Arab berharakat (fathah, kasrah, dammah, sukun, syaddah, tanwin).
ARABIC_QUESTION = "إِلَى أَيْنَ ذَهَبَ أَحْمَدُ؟"
ARABIC_STIMULUS = "ذَهَبَ أَحْمَدُ إِلَى الْمَدْرَسَةِ صَبَاحًا."
ARABIC_CHOICES = ["إِلَى الْمَدْرَسَةِ", "إِلَى السُّوْقِ", "إِلَى الْبَيْتِ"]
MATH_QUESTION = r"Nilai dari $\frac{1}{2} + \frac{1}{4}$ adalah ..."
MATH_DISPLAY = r"$$\begin{cases} x + y = 5 \\ x - y = 1 \end{cases}$$"


def harakat(text):
    """Jumlah tanda harakat (karakter penanda Unicode kategori Mn)."""
    return sum(1 for c in text if unicodedata.category(c) == "Mn")


def excel(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Soal"
    sheet.append([f"{c.label}*" if c.required else c.label for c in COLUMNS])
    for values in rows:
        sheet.append([values.get(c.key, "") for c in COLUMNS])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile("soal.xlsx", buffer.getvalue())


def zip_of(files):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return SimpleUploadedFile("gambar.zip", buffer.getvalue())


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class QuestionImportTests(AssessmentTestBase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(TEMP_MEDIA, ignore_errors=True)

    def setUp(self):
        super().setUp()
        self.assessment = self.make_assessment(published=False)
        self.url = reverse("assessment-import-questions", args=[self.assessment.id])
        self.client.force_authenticate(self.teacher)

    def upload(self, dry_run=False, **data):
        return self.client.post(self.url, {"dry_run": str(dry_run).lower(), **data}, format="multipart")

    def imported(self):
        # Dua soal pertama berasal dari make_assessment(); soal impor ada sesudahnya.
        return list(self.assessment.questions.order_by("order")[2:])

    # ------------------------------------------------------------ template

    def test_template_has_cheatsheet_and_examples(self):
        res = self.client.get(reverse("assessment-question-template"))
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        workbook = load_workbook(io.BytesIO(res.content))
        self.assertEqual(workbook.active.title, "Soal")
        self.assertIn("Contekan Rumus", workbook.sheetnames)
        cheats = [row[1] for row in workbook["Contekan Rumus"].iter_rows(min_row=2, values_only=True)]
        for latex in (r"\frac", r"\sqrt", r"\log", r"\sin", r"\int", r"\begin{pmatrix}", r"\begin{cases}"):
            self.assertTrue(any(latex in c for c in cheats), latex)
        guide = "\n".join(
            str(v) for row in workbook["Petunjuk"].iter_rows(values_only=True) for v in row if v
        )
        self.assertIn(r"\frac{1}{2}", guide)  # contoh soal matematika
        self.assertIn("ذَهَبَ أَحْمَدُ", guide)  # contoh soal bahasa Arab (berharakat)
        self.assertIn("rangkaian1.png", guide)  # contoh kolom gambar

    def test_filled_template_can_be_imported_back(self):
        res = self.client.get(reverse("assessment-question-template"))
        workbook = load_workbook(io.BytesIO(res.content))
        sheet = workbook["Soal"]
        sheet.append(["pg", "", "", MATH_QUESTION, "", "otomatis", 2, r"$\frac{3}{4}$", r"$\frac{2}{6}$"]
                     + [""] * 8 + ["A"])  # fmt: skip
        buffer = io.BytesIO()
        workbook.save(buffer)
        res = self.upload(file=SimpleUploadedFile("isi.xlsx", buffer.getvalue()))
        self.assertEqual(res.status_code, status.HTTP_201_CREATED, res.data)
        self.assertEqual(self.imported()[0].text, MATH_QUESTION)

    # ------------------------------------------------------------ rumus

    def test_excel_latex_preserved_exactly(self):
        rows = [
            {"kind": "pg", "text": MATH_QUESTION, "choice_A": r"$\frac{3}{4}$", "choice_B": r"$\sqrt[3]{27}$",
             "choice_C": "1/2", "key": "A", "points": "2"},
            {"kind": "esai", "text": "Selesaikan " + MATH_DISPLAY},
        ]  # fmt: skip
        res = self.upload(file=excel(rows))
        self.assertEqual(res.status_code, status.HTTP_201_CREATED, res.data)
        mc, essay = self.imported()
        self.assertEqual(mc.text, MATH_QUESTION)
        self.assertEqual(mc.choices, [r"$\frac{3}{4}$", r"$\sqrt[3]{27}$", "1/2"])
        self.assertEqual((mc.correct_choice, mc.points), (0, 2))
        self.assertEqual(essay.text, "Selesaikan " + MATH_DISPLAY)
        self.assertEqual((essay.kind, essay.points), ("essay", 5))

    def test_paste_latex_keeps_display_math_on_its_own_line(self):
        text = f"1. {MATH_QUESTION}\nA. $\\frac{{3}}{{4}}$\n*B. $\\frac{{2}}{{6}}$\n\n2. [esai] Tentukan $x$ dari\n{MATH_DISPLAY}\n"
        res = self.upload(mode="text", text=text)
        self.assertEqual(res.status_code, status.HTTP_201_CREATED, res.data)
        mc, essay = self.imported()
        self.assertEqual(mc.choices, [r"$\frac{3}{4}$", r"$\frac{2}{6}$"])
        self.assertEqual(mc.correct_choice, 1)  # tanda * = kunci
        self.assertEqual(essay.text, "Tentukan $x$ dari\n" + MATH_DISPLAY)

    # ------------------------------------------------------------ bahasa Arab

    def test_excel_arabic_harakat_preserved(self):
        row = {"kind": "pg", "stimulus": ARABIC_STIMULUS, "text": ARABIC_QUESTION, "direction": "rtl",
               "choice_A": ARABIC_CHOICES[0], "choice_B": ARABIC_CHOICES[1], "choice_C": ARABIC_CHOICES[2],
               "key": "ب"}  # fmt: skip
        res = self.upload(file=excel([row]))
        self.assertEqual(res.status_code, status.HTTP_201_CREATED, res.data)
        question = self.imported()[0]
        self.assertEqual(question.text, ARABIC_QUESTION)
        self.assertEqual(harakat(question.text), harakat(ARABIC_QUESTION))
        self.assertEqual(question.stimulus, ARABIC_STIMULUS)
        self.assertEqual(question.choices, ARABIC_CHOICES)
        self.assertEqual((question.text_direction, question.correct_choice), ("rtl", 1))  # kunci huruf Arab

    def test_paste_from_word_arabic_numbering_and_hidden_marks(self):
        # Salinan dari Word sering membawa tanda RLM/LRM tak terlihat, NBSP, dan akhir baris CRLF.
        text = (
            f"Stimulus: {RLM}{ARABIC_STIMULUS}\r\n"
            f"{RLM}١. {ARABIC_QUESTION}\r\n"
            f"{RLM}أ. {ARABIC_CHOICES[0]}\r\n"
            f"{RLM}ب.{NBSP}{ARABIC_CHOICES[1]}{LRM}\r\n"
            f"{RLM}ج. {ARABIC_CHOICES[2]}\r\n"
            "Kunci: أ\r\nArah: rtl\r\n"
            "2. Arti kata «الْمَدْرَسَةُ» adalah ...\r\nA. Sekolah\r\nB. Masjid\r\nKunci: A\r\n"
        )
        res = self.upload(mode="text", text=text)
        self.assertEqual(res.status_code, status.HTTP_201_CREATED, res.data)
        arabic, mixed = self.imported()
        self.assertEqual(arabic.text, ARABIC_QUESTION)
        self.assertEqual(arabic.stimulus, ARABIC_STIMULUS)
        self.assertEqual(arabic.choices, ARABIC_CHOICES)  # tanda tak terlihat & NBSP dibersihkan
        self.assertEqual(sum(harakat(c) for c in arabic.choices), sum(harakat(c) for c in ARABIC_CHOICES))
        self.assertEqual((arabic.correct_choice, arabic.text_direction), (0, "rtl"))
        self.assertIn("الْمَدْرَسَةُ", mixed.text)  # teks campuran Indonesia-Arab utuh
        self.assertEqual(mixed.text_direction, "auto")

    def test_parse_text_ignores_title_and_reads_points(self):
        items = parse_text("LATIHAN SOAL BAB 1\n\n1. Soal pertama\nA. satu\nB. dua\nKunci: B\nPoin: 3")
        self.assertEqual(len(items), 1)
        data = items[0]["data"]
        self.assertEqual((data["correct_choice"], data["points"], data["choices"]), (1, 3, ["satu", "dua"]))

    # ------------------------------------------------------------ gambar

    def test_excel_with_zip_images(self):
        row = {"kind": "pg", "stimulus": "Perhatikan grafik.", "stimulus_image": "grafik.PNG",
               "text": "Grafik mana yang naik?", "question_image": "soal1.jpg",
               "choice_A": "", "image_A": "a.png", "choice_B": "Grafik B", "image_B": "b.png", "key": "A"}  # fmt: skip
        images = zip_of(
            {"Gambar/grafik.png": PNG, "soal1.jpg": JPG, "a.png": PNG, "b.png": PNG, "__MACOSX/._a.png": b"x"}
        )
        res = self.upload(file=excel([row]), images=images)
        self.assertEqual(res.status_code, status.HTTP_201_CREATED, res.data)
        question = self.imported()[0]
        self.assertEqual(question.choices, ["", "Grafik B"])  # pilihan A hanya berupa gambar
        slots = {(i.slot, i.index): i.original_filename for i in question.images.all()}
        self.assertEqual(
            slots,
            {
                ("stimulus", 0): "grafik.png",
                ("question", 0): "soal1.jpg",
                ("choice", 0): "a.png",
                ("choice", 1): "b.png",
            },
        )

    def test_paste_with_zip_images(self):
        text = "Gambar stimulus: peta.png\n1. Pulau terbesar?\nGambar: soal.png\nA. Jawa [gambar: jawa.png]\nB. Kalimantan\nKunci: B"
        images = zip_of({"peta.png": PNG, "soal.png": PNG, "jawa.png": PNG})
        res = self.upload(mode="text", text=text, images=images)
        self.assertEqual(res.status_code, status.HTTP_201_CREATED, res.data)
        question = self.imported()[0]
        self.assertEqual(question.choices, ["Jawa", "Kalimantan"])
        self.assertEqual(question.images.count(), 3)

    def test_missing_or_invalid_images_rejected_and_nothing_saved(self):
        row = {
            "kind": "pg",
            "text": "Soal",
            "question_image": "tidak-ada.png",
            "choice_A": "a",
            "choice_B": "b",
            "key": "A",
        }
        res = self.upload(file=excel([row]))
        self.assertIn("unggah juga file ZIP", " ".join(res.data["rows"][0]["errors"]))
        res = self.upload(file=excel([row]), images=zip_of({"lain.png": PNG}))
        self.assertIn("tidak ada di dalam file ZIP", " ".join(res.data["rows"][0]["errors"]))
        res = self.upload(file=excel([row]), images=zip_of({"tidak-ada.png": b"bukan gambar"}))
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("bukan gambar yang valid", str(res.data["images"]))
        res = self.upload(file=excel([row]), images=zip_of({"tidak-ada.svg": b"<svg/>"}))
        self.assertIn("PNG, JPG", str(res.data["images"]))
        self.assertEqual(len(self.imported()), 0)
        self.assertEqual(QuestionImage.objects.count(), 0)

    # ------------------------------------------------------------ aturan umum

    def test_dry_run_and_all_or_nothing(self):
        rows = [
            {"kind": "pg", "text": "Benar", "choice_A": "a", "choice_B": "b", "key": "A"},
            {"kind": "pg", "text": "Tanpa kunci", "choice_A": "a", "choice_B": "b"},
            {"kind": "kuis", "text": "Jenis salah"},
        ]
        res = self.upload(dry_run=True, file=excel(rows))
        self.assertEqual(res.data["summary"], {"total": 3, "valid": 1, "invalid": 2})
        self.assertIn("Kunci", " ".join(res.data["rows"][1]["errors"]))
        self.assertIn("Jenis", " ".join(res.data["rows"][2]["errors"]))
        self.assertEqual(self.upload(file=excel(rows)).status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(len(self.imported()), 0)

    def test_locked_after_attempt_and_only_owner_teacher(self):
        self.client.force_authenticate(self.other_teacher)
        row = {"kind": "esai", "text": "x"}
        self.assertEqual(self.upload(file=excel([row])).status_code, status.HTTP_404_NOT_FOUND)
        self.act(self.student, "start", self.make_assessment())
        locked = self.make_assessment()
        self.act(self.student2, "start", locked)
        self.client.force_authenticate(self.teacher)
        res = self.client.post(
            reverse("assessment-import-questions", args=[locked.id]), {"file": excel([row])}
        )
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class QuestionImageAccessTests(AssessmentTestBase):
    def setUp(self):
        super().setUp()
        self.assessment = self.make_assessment()
        self.client.force_authenticate(self.teacher)
        url = reverse("assessment-question-images", args=[self.mc.id])
        res = self.client.post(url, {"slot": "question", "file": SimpleUploadedFile("soal.png", PNG)})
        self.assertEqual(res.status_code, status.HTTP_200_OK, res.data)
        self.image_url = "/api" + res.data["images"]["question"]["url"]

    def test_teacher_can_view_student_only_after_starting(self):
        self.assertEqual(self.client.get(self.image_url).status_code, status.HTTP_200_OK)
        self.client.force_authenticate(self.student)
        self.assertEqual(self.client.get(self.image_url).status_code, status.HTTP_404_NOT_FOUND)
        questions = self.act(self.student, "start", self.assessment).data["questions"]
        self.assertIsNotNone(questions[0]["images"]["question"])
        self.assertNotIn("correct_choice", questions[0])
        self.client.force_authenticate(self.student)
        res = self.client.get(self.image_url)
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(b"".join(res.streaming_content), PNG)
        self.client.force_authenticate(self.student_tkj)
        self.assertEqual(self.client.get(self.image_url).status_code, status.HTTP_404_NOT_FOUND)

    def test_upload_validation_delete_and_lock(self):
        url = reverse("assessment-question-images", args=[self.mc.id])
        res = self.client.post(url, {"slot": "choice", "index": 9, "file": SimpleUploadedFile("a.png", PNG)})
        self.assertIn("index", res.data)
        res = self.client.post(url, {"slot": "stimulus", "file": SimpleUploadedFile("x.png", b"palsu")})
        self.assertIn("file", res.data)
        self.assertEqual(self.client.delete(url + "?slot=question").status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(QuestionImage.objects.exists())
        self.act(self.student, "start", self.assessment)
        self.client.force_authenticate(self.teacher)
        res = self.client.post(url, {"slot": "question", "file": SimpleUploadedFile("s.png", PNG)})
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)

    def test_image_files_removed_when_question_deleted(self):
        image = QuestionImage.objects.get()
        path = image.file.path
        Question.objects.filter(pk=self.mc.pk).delete()  # hapus lewat queryset (tanpa delete() per objek)
        import os

        self.assertFalse(os.path.exists(path))
