"""Helper bersama untuk template & impor Excel (.xlsx).

Dipakai impor mapel, kelas, pengampu, dan soal asesmen. Polanya sama dengan impor pengguna
(apps/accounts/user_import.py): unduh template -> isi -> periksa (tidak ada yang disimpan) ->
impor jika SEMUA baris benar.
"""

import io
import os
from dataclasses import dataclass

from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from rest_framework.exceptions import ValidationError

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DEFAULT_MAX_ROWS = 500
HEADER_FILL = PatternFill("solid", fgColor="1F8A4C")


@dataclass
class Column:
    key: str
    label: str
    required: bool = False
    width: int = 18
    text: bool = False  # format teks agar angka seperti NIP/kode tidak berubah
    # Pilihan dropdown: daftar nilai (disimpan di lembar tersembunyi agar tidak dibatasi panjang).
    choices: list | None = None


def normalize(text):
    return " ".join(str(text).split()).lower()


def cell_text(value):
    """Isi sel menjadi teks. Angka seperti 26001 (terbaca 26001.0) menjadi '26001'."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def build_template(
    sheet_title, columns, guide_title, guide_lines, examples, max_rows=DEFAULT_MAX_ROWS, extra_sheets=()
):
    """Buat file template: lembar Petunjuk (dengan contoh) + lembar isian dengan dropdown.
    `extra_sheets`: [(judul, baris_pertama_judul_kolom, [baris...])], mis. lembar Contekan Rumus."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    lists = None
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=f"{column.label}*" if column.required else column.label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        letter = cell.column_letter
        sheet.column_dimensions[letter].width = column.width
        if column.text:
            for row in range(2, max_rows + 2):
                sheet[f"{letter}{row}"].number_format = "@"
        if column.choices:
            if lists is None:
                lists = workbook.create_sheet("Daftar")
                lists.sheet_state = "hidden"
            for row, value in enumerate(column.choices, start=1):
                lists.cell(row=row, column=index, value=value)
            validation = DataValidation(
                type="list",
                formula1=f"Daftar!${letter}$1:${letter}${len(column.choices)}",
                allow_blank=not column.required,
            )
            validation.errorStyle = "warning"
            validation.error = "Nilai tidak ada di daftar. Periksa kembali."
            sheet.add_data_validation(validation)
            validation.add(f"{letter}2:{letter}{max_rows + 1}")
    sheet.freeze_panes = "A2"

    guide = workbook.create_sheet("Petunjuk", 0)
    guide.cell(row=1, column=1, value=guide_title).font = Font(bold=True, size=13)
    lines = [
        "",
        f"Isi data di lembar '{sheet_title}' mulai baris ke-2. Maksimal {max_rows} baris per file.",
        "Kolom bertanda * wajib diisi. Kolom berdropdown: pilih dari daftar.",
        *guide_lines,
        "",
        "Contoh isian:",
    ]
    for row, text in enumerate(lines, start=2):
        guide.cell(row=row, column=1, value=text).font = Font(bold=text == "Contoh isian:")
    start = len(lines) + 2
    for offset, values in enumerate([[c.label for c in columns], *examples]):
        for column, value in enumerate(values, start=1):
            guide.cell(row=start + offset, column=column, value=value).font = Font(bold=offset == 0)
    guide.column_dimensions["A"].width = 30

    for title, header, rows in extra_sheets:
        extra = workbook.create_sheet(title)
        extra.append(header)
        for cell in extra[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL
        for row in rows:
            extra.append(row)
        for letter, width in zip("ABCD", (34, 46, 40, 30)):
            extra.column_dimensions[letter].width = width
        for row in extra.iter_rows(min_row=2):
            for cell in row:
                # Teks apa adanya: rumus LaTeX (mis. rac) tidak diubah Excel.
                cell.number_format = "@"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        extra.freeze_panes = "A2"

    workbook.active = workbook.index(sheet)
    sheet.sheet_view.tabSelected = True
    guide.sheet_view.tabSelected = False

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def template_response(content, filename):
    response = HttpResponse(content, content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def read_rows(file, sheet_title, columns, max_rows=DEFAULT_MAX_ROWS):
    """Baca lembar isian. Mengembalikan [{"row": nomor_baris, "data": {key: teks}}]."""
    if file is None:
        raise ValidationError({"file": "Pilih file Excel yang akan diimpor."})
    if os.path.splitext(file.name)[1].lower() != ".xlsx":
        raise ValidationError(
            {"file": "Gunakan file Excel berformat .xlsx (unduh template yang disediakan)."}
        )
    try:
        workbook = load_workbook(file, read_only=True, data_only=True)
    except Exception:
        raise ValidationError({"file": "File tidak bisa dibaca. Pastikan file Excel .xlsx yang valid."})

    sheet = workbook[sheet_title] if sheet_title in workbook.sheetnames else workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None) or ()
    labels = {normalize(c.label): c.key for c in columns}
    positions = {}
    for index, title in enumerate(header):
        key = labels.get(normalize(cell_text(title).replace("*", "")))
        if key:
            positions[key] = index
    missing = [c.label for c in columns if c.required and c.key not in positions]
    if missing:
        raise ValidationError(
            {"file": f"Kolom {', '.join(missing)} tidak ditemukan. Gunakan template yang disediakan."}
        )

    result = []
    for row_number, values in enumerate(rows, start=2):
        data = {
            key: cell_text(values[index]) if index < len(values) else "" for key, index in positions.items()
        }
        if not any(data.values()):
            continue  # baris kosong dilewati
        result.append({"row": row_number, "data": {c.key: data.get(c.key, "") for c in columns}})
        if len(result) > max_rows:
            raise ValidationError(
                {"file": f"Maksimal {max_rows} baris per file. Bagi menjadi beberapa file."}
            )
    workbook.close()
    if not result:
        raise ValidationError({"file": f"Tidak ada data di lembar '{sheet_title}'."})
    return result


def first_errors(errors, labels=None):
    """Ubah error serializer DRF {field: [pesan]} menjadi daftar teks sederhana.
    `labels`: nama kolom yang ditampilkan, mis. {"correct_choice": "Kunci"}."""
    labels = labels or {}
    messages = []
    for field, value in errors.items():
        items = value if isinstance(value, list) else [value]
        for item in items:
            text = str(item) if not isinstance(item, dict) else " ".join(map(str, item.values()))
            messages.append(
                text if field in ("non_field_errors", "detail") else f"{labels.get(field, field)}: {text}"
            )
    return messages


def import_response(rows, dry_run, save):
    """Respons standar impor. `rows`: [{"row", ..., "errors": []}]. `save()` dipanggil hanya jika
    dry_run=False dan SEMUA baris benar; hasilnya (jumlah yang disimpan) dikembalikan."""
    from rest_framework import status
    from rest_framework.response import Response

    invalid = [r for r in rows if r["errors"]]
    summary = {"total": len(rows), "valid": len(rows) - len(invalid), "invalid": len(invalid)}
    if dry_run:
        return Response({"imported": False, "summary": summary, "rows": rows})
    if invalid:
        return Response(
            {
                "detail": f"{len(invalid)} baris masih salah. Perbaiki lalu coba lagi.",
                "summary": summary,
                "rows": rows,
            },
            status=status.HTTP_400_BAD_REQUEST,
        )
    return Response({"imported": True, "summary": summary, "created": save()}, status=status.HTTP_201_CREATED)


def is_dry_run(request):
    return str(request.data.get("dry_run", "true")).lower() != "false"
